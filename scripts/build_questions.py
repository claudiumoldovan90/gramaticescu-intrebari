#!/usr/bin/env python3
"""
Gramaticescu (repository-ul public cu întrebări) – transformă questions.xlsx în fișierele
publicate pe GitHub Pages, de unde le citește aplicația.

Ce face:
  1. citește foile „Întrebări” și „Reguli” din questions.xlsx;
  2. verifică întrebările (oprește publicarea la erori și arată rândul greșit);
  3. pune data de activare pentru întrebările noi (data/activation.json);
  4. scrie data/questions.json (pentru istoricul din git) și folderul site/:
       site/questions.json  – întrebările (citite de aplicație)
       site/manifest.json   – versiunea (fișier mic, verificat de aplicație la pornire)
       site/index.html      – o pagină simplă de control: versiunea și numărul de întrebări

Rulare locală (opțional):  python3 scripts/build_questions.py
"""
import datetime as dt
import hashlib
import json
import os
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Lipsește openpyxl. Instalează cu: pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "questions.xlsx")
DATA_DIR = os.path.join(ROOT, "data")
SITE_DIR = os.path.join(ROOT, "site")
ACTIVATION = os.path.join(DATA_DIR, "activation.json")

ACTIVATION_DELAY_DAYS = 7   # întrebările noi intră în provocarea zilei după atâtea zile
MAX_PAIR_LETTERS = 26       # peste atâtea litere (corect + greșit) primești un avertisment
SHEET_Q, SHEET_R = "Întrebări", "Reguli"
COLUMNS = ["ID", "Temă", "Întrebarea", "Răspuns corect", "Răspuns greșit", "Activă"]

IN_ACTIONS = os.environ.get("GITHUB_ACTIONS") == "true"
errors, warnings = [], []


def err(row, msg):
    text = f"Foaia „{SHEET_Q}”, rândul {row}: {msg}" if row else msg
    errors.append(text)
    print(f"::error title=Eroare în questions.xlsx::{text}" if IN_ACTIONS else f"EROARE  {text}")


def warn(row, msg):
    text = f"Foaia „{SHEET_Q}”, rândul {row}: {msg}"
    warnings.append(text)
    print(f"::warning title=Atenție în questions.xlsx::{text}" if IN_ACTIONS else f"ATENȚIE {text}")


def cell(v):
    return "" if v is None else str(v).strip()


def today_bucharest():
    try:
        from zoneinfo import ZoneInfo
        return dt.datetime.now(ZoneInfo("Europe/Bucharest")).date()
    except Exception:
        return (dt.datetime.utcnow() + dt.timedelta(hours=3)).date()


def read_workbook():
    if not os.path.exists(XLSX):
        sys.exit("Nu găsesc questions.xlsx în rădăcina repository-ului.")
    wb = load_workbook(XLSX, data_only=True)
    for name in (SHEET_Q, SHEET_R):
        if name not in wb.sheetnames:
            sys.exit(f"Lipsește foaia „{name}” din questions.xlsx.")

    # reguli
    themes, seen_t = [], set()
    for r, row in enumerate(wb[SHEET_R].iter_rows(min_row=2, values_only=True), start=2):
        name, rule = cell(row[0] if row else None), cell(row[1] if len(row) > 1 else None)
        if not name and not rule:
            continue
        if not name or not rule:
            err(None, f"Foaia „{SHEET_R}”, rândul {r}: tema și regula trebuie completate amândouă.")
            continue
        if name in seen_t:
            err(None, f"Foaia „{SHEET_R}”, rândul {r}: tema „{name}” apare de două ori.")
            continue
        seen_t.add(name)
        themes.append({"name": name, "rule": rule.replace("\r\n", "\n")})

    # întrebări
    ws = wb[SHEET_Q]
    header = [cell(c.value) for c in ws[1]]
    if header[: len(COLUMNS)] != COLUMNS:
        sys.exit(f"Capul de tabel din „{SHEET_Q}” trebuie să fie: {', '.join(COLUMNS)}. Am găsit: {', '.join(header)}")

    questions, ids = [], {}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals = [cell(v) for v in (list(row) + [None] * 6)[:6]]
        if not any(vals):
            continue
        sid, theme, text, correct, wrong, active = vals
        try:
            qid = int(float(sid))
            if qid <= 0 or float(sid) != qid:
                raise ValueError
        except ValueError:
            err(r, f"ID-ul „{sid}” trebuie să fie un număr întreg pozitiv.")
            continue
        if qid in ids:
            err(r, f"ID-ul {qid} apare și pe rândul {ids[qid]}. Fiecare întrebare are nevoie de un ID unic.")
            continue
        ids[qid] = r
        if theme not in seen_t:
            err(r, f"tema „{theme}” nu există în foaia „{SHEET_R}”.")
        if text.count("___") != 1:
            err(r, "întrebarea trebuie să conțină exact un spațiu gol scris ___ (trei liniuțe jos).")
        if not correct or not wrong:
            err(r, "completează și răspunsul corect, și răspunsul greșit.")
        elif correct == wrong:
            err(r, f"răspunsul corect și cel greșit sunt identice („{correct}”).")
        a = active.lower()
        if a not in ("", "da", "nu"):
            err(r, f"coloana Activă trebuie să fie „da” sau „nu” (am găsit „{active}”).")
        if correct and wrong and len(correct) + len(wrong) > MAX_PAIR_LETTERS:
            warn(r, f"variantele „{correct}” / „{wrong}” sunt lungi și pe telefon vor apărea una sub alta.")
        questions.append({"id": qid, "theme": theme, "text": text, "correct": correct,
                          "wrong": wrong, "active": a != "nu"})
    if not questions:
        err(None, "Nu am găsit nicio întrebare.")
    questions.sort(key=lambda q: q["id"])
    return themes, questions


def apply_activation(questions):
    """Întrebările noi primesc o dată de activare, ca toți jucătorii să aibă lista nouă
    înainte ca ele să apară în provocarea zilei. La prima rulare, toate sunt active imediat."""
    first_run = not os.path.exists(ACTIVATION)
    known = {}
    if not first_run:
        with open(ACTIVATION, encoding="utf-8") as fh:
            known = {int(k): v for k, v in json.load(fh).items()}
    start = (today_bucharest() + dt.timedelta(days=ACTIVATION_DELAY_DAYS)).isoformat()
    for q in questions:
        if q["id"] not in known:
            known[q["id"]] = None if first_run else start
        if known[q["id"]]:
            q["activeFrom"] = known[q["id"]]
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(ACTIVATION, "w", encoding="utf-8") as fh:
        json.dump({str(k): known[k] for k in sorted(known)}, fh, ensure_ascii=False, indent=1)


def main():
    themes, questions = read_workbook()
    if errors:
        print(f"\n{len(errors)} erori. Publicarea s-a oprit; jucătorii păstrează lista anterioară.")
        sys.exit(1)
    apply_activation(questions)

    content = {"themes": themes, "questions": questions}
    digest = hashlib.sha256(json.dumps(content, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    version = dt.datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    data = {"version": version, "generated": dt.datetime.utcnow().isoformat(timespec="seconds") + "Z", **content}
    manifest = {"version": version, "count": len(questions),
                "active": sum(q["active"] for q in questions), "sha256": digest}

    # copie lizibilă în repository (fără versiune, ca istoricul să arate doar schimbările reale)
    with open(os.path.join(DATA_DIR, "questions.json"), "w", encoding="utf-8") as fh:
        json.dump(content, fh, ensure_ascii=False, indent=1)

    os.makedirs(SITE_DIR, exist_ok=True)
    with open(os.path.join(SITE_DIR, "questions.json"), "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(SITE_DIR, "manifest.json"), "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False)

    # pagină de control, ca să vezi rapid ce versiune e publicată
    themes_count = {}
    for q in questions:
        themes_count[q["theme"]] = themes_count.get(q["theme"], 0) + 1
    rows = "".join(f"<tr><td>{t['name']}</td><td>{themes_count.get(t['name'], 0)}</td></tr>" for t in themes)
    page = f"""<!doctype html><html lang="ro"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Întrebări Gramaticescu</title><style>body{{font-family:system-ui,sans-serif;max-width:640px;margin:0 auto;padding:24px 16px;color:#1f2a44;background:#fdfdf8}}
table{{border-collapse:collapse;width:100%}}td,th{{border-bottom:1px solid #d6e4f5;padding:6px 4px;text-align:left}}code{{background:#e3ecfb;padding:1px 5px;border-radius:4px}}</style></head>
<body><h1>Întrebări Gramaticescu</h1><p>Versiunea <code>{version}</code> · {len(questions)} întrebări ({manifest['active']} active) · {len(themes)} teme</p>
<p>Fișiere pentru aplicație: <a href="manifest.json">manifest.json</a> · <a href="questions.json">questions.json</a></p>
<table><tr><th>Temă</th><th>Întrebări</th></tr>{rows}</table></body></html>"""
    with open(os.path.join(SITE_DIR, "index.html"), "w", encoding="utf-8") as fh:
        fh.write(page)

    active = manifest["active"]
    print(f"Gata: {len(questions)} întrebări ({active} active), {len(themes)} teme, versiunea {version}.")
    if warnings:
        print(f"{len(warnings)} avertismente (nu opresc publicarea).")
    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary:
        with open(summary, "a", encoding="utf-8") as fh:
            fh.write(f"### Întrebări publicate\n\n- Versiunea: `{version}`\n- Întrebări: {len(questions)} ({active} active)\n- Teme: {len(themes)}\n")
            if warnings:
                fh.write("\n**Avertismente:**\n\n" + "\n".join(f"- {w}" for w in warnings) + "\n")


if __name__ == "__main__":
    main()
